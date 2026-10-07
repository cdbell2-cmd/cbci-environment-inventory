#!/usr/bin/env python

# pip install openpyxl

"""
Turn a timestamped environment-inventory run into clean, client-friendly Excel workbooks.

Produces, under <run>/reports/:
  - 00_Environment_Inventory_Summary.xlsx   one workbook summarizing the whole environment
  - operations-center.xlsx                  the Operations Center (CJOC)
  - controller__<name>.xlsx                 one workbook per controller

Each workbook is organized around the three CIBC Success Path objectives:
  1. Controller resource allocation
  2. Job ecosystem
  3. Plugin usage & health

Usage:
  python build_excel_reports.py                     # newest run under ./output
  python build_excel_reports.py --run-dir output/20261007-113134_environment-inventory
  python build_excel_reports.py --run-dir <dir> --output-dir <dir>/reports
"""

import argparse
import csv
import json
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl is required.  Install it with:  pip install openpyxl")
    sys.exit(1)

# ---- palette -------------------------------------------------------------------------
HEADER_FILL = PatternFill('solid', fgColor='1F3864')      # dark blue
HEADER_FONT = Font(bold=True, color='FFFFFF', size=11)
TITLE_FONT = Font(bold=True, size=16, color='1F3864')
SUBTITLE_FONT = Font(bold=True, size=12, color='1F3864')
LABEL_FONT = Font(bold=True)
FREESTYLE_FILL = PatternFill('solid', fgColor='FCE4D6')   # light amber: Pipeline-coaching target
FLAG_FILL = PatternFill('solid', fgColor='FFF2CC')        # light yellow: flagged plugin
WRAP = Alignment(wrap_text=True, vertical='top')
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ---- small IO helpers ----------------------------------------------------------------
def read_csv(path):
    """Return (header, rows) from a CSV file, or (None, []) if missing/empty."""
    p = Path(path)
    if not p.is_file():
        return None, []
    with open(p, 'r', newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        all_rows = list(reader)
    if not all_rows:
        return None, []
    return all_rows[0], all_rows[1:]


def read_json(path):
    p = Path(path)
    if not p.is_file():
        return None
    try:
        with open(p, 'r', encoding='utf-8-sig') as f:
            return json.load(f)
    except Exception:
        return None


# ---- sheet builders ------------------------------------------------------------------
def _autosize(ws, max_width=70):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            col = cell.column_letter
            longest = max((len(line) for line in str(cell.value).splitlines()), default=0)
            widths[col] = max(widths.get(col, 0), longest)
    for col, w in widths.items():
        ws.column_dimensions[col].width = min(max(w + 2, 10), max_width)


def write_table(ws, header, rows, start_row=1, highlight=None):
    """Write a styled table (bold header, filter, frozen header) starting at start_row."""
    if not header:
        ws.cell(row=start_row, column=1, value='(no data)')
        return start_row + 1
    for c, name in enumerate(header, start=1):
        cell = ws.cell(row=start_row, column=c, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical='center')
        cell.border = BORDER
    r = start_row + 1
    for row in rows:
        fill = highlight(header, row) if highlight else None
        for c, _name in enumerate(header, start=1):
            val = row[c - 1] if c - 1 < len(row) else ''
            cell = ws.cell(row=r, column=c, value=_coerce(val))
            cell.border = BORDER
            if fill:
                cell.fill = fill
        r += 1
    last_col = get_column_letter(len(header))
    ws.auto_filter.ref = f"A{start_row}:{last_col}{max(r - 1, start_row)}"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    _autosize(ws)
    return r


def _coerce(val):
    """Turn numeric-looking strings into numbers so Excel sorts/sums them correctly."""
    if val is None:
        return ''
    s = str(val)
    if s == '' or s in ('N/A', 'none'):
        return s
    try:
        if s.lstrip('-').isdigit():
            return int(s)
        f = float(s)
        return f
    except ValueError:
        return s


def kv_sheet(ws, title, pairs, start_row=1):
    """Write a two-column Field/Value block with a title above it."""
    ws.cell(row=start_row, column=1, value=title).font = SUBTITLE_FONT
    r = start_row + 1
    for label, value in pairs:
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = LABEL_FONT
        lc.border = BORDER
        vc = ws.cell(row=r, column=2, value=_coerce(value))
        vc.border = BORDER
        vc.alignment = WRAP
        r += 1
    _autosize(ws)
    return r + 1


def highlight_freestyle(header, row):
    try:
        return FREESTYLE_FILL if row[header.index('Type')].strip().lower() == 'freestyle' else None
    except (ValueError, IndexError):
        return None


# ---- friendly labels for resources.json ----------------------------------------------
RES_LABELS = [
    ('name', 'Controller Name'), ('team', 'Owning Team'), ('url', 'URL'),
    ('version', 'CloudBees CI Version'), ('cpus', 'CPU (cores)'), ('memory_mb', 'Memory (MB)'),
    ('disk_gb', 'Disk (GB)'), ('ha', 'High Availability'), ('ha_inferred', 'HA Inferred'),
    ('replicas', 'Replicas'), ('state', 'State'), ('online', 'Online'),
    ('className', 'Controller Type'), ('provisioning_source_class', 'Provisioning Source'),
]


# ======================================================================================
# Per-controller workbook
# ======================================================================================
def build_controller_workbook(folder, out_path):
    res = read_json(folder / 'resources.json') or {}
    jt_header, jt_rows = read_csv(folder / 'job-type-summary.csv')
    jobs_header, jobs_rows = read_csv(folder / 'jobs.csv')
    ag_header, ag_rows = read_csv(folder / 'static-agents.csv')
    pl_header, pl_rows = read_csv(folder / 'plugins.csv')
    ph_header, ph_rows = read_csv(folder / 'plugin-health-report.csv')

    wb = Workbook()

    # -- Overview ----------------------------------------------------------------------
    ws = wb.active
    ws.title = 'Overview'
    ws.cell(row=1, column=1, value=f"Controller: {res.get('name', folder.name)}").font = TITLE_FONT
    job_counts = {r[0]: r[1] for r in jt_rows if len(r) >= 2}
    total_jobs = sum(int(v) for k, v in job_counts.items() if k != 'Folder' and str(v).isdigit())
    overview = [
        ('Controller Name', res.get('name')),
        ('Owning Team', res.get('team')),
        ('URL', res.get('url')),
        ('CloudBees CI Version', res.get('version')),
        ('CPU (cores)', res.get('cpus')),
        ('Memory (MB)', res.get('memory_mb')),
        ('High Availability', res.get('ha')),
        ('Replicas', res.get('replicas')),
        ('State', res.get('state')),
        ('Total Jobs (excl. folders)', total_jobs),
        ('Freestyle Jobs (coaching targets)', job_counts.get('Freestyle', 0)),
        ('Installed Plugins', len(pl_rows)),
        ('Static Agents', len(ag_rows)),
    ]
    kv_sheet(ws, 'Key facts', overview, start_row=3)

    # -- 1. Resources & Agents ---------------------------------------------------------
    ws = wb.create_sheet('Resources & Agents')
    ws.cell(row=1, column=1, value='Objective 1 — Controller Resource Allocation').font = TITLE_FONT
    pairs = [(label, res.get(key)) for key, label in RES_LABELS if res.get(key) is not None]
    next_row = kv_sheet(ws, 'Resource allocation', pairs, start_row=3)
    ws.cell(row=next_row, column=1, value='Static (non-Kubernetes) Agents').font = SUBTITLE_FONT
    if ag_header:
        write_table(ws, ag_header, ag_rows, start_row=next_row + 1)
    else:
        ws.cell(row=next_row + 1, column=1, value='None')

    # -- 2. Jobs -----------------------------------------------------------------------
    ws = wb.create_sheet('Job Types')
    ws.cell(row=1, column=1, value='Objective 2 — Job Ecosystem (type summary)').font = TITLE_FONT
    if jt_header:
        write_table(ws, jt_header, jt_rows, start_row=3)

    ws = wb.create_sheet('Jobs')
    ws.cell(row=1, column=1, value='Objective 2 — Jobs (amber = Freestyle, Pipeline-coaching targets)').font = TITLE_FONT
    if jobs_header:
        write_table(ws, jobs_header, jobs_rows, start_row=3, highlight=highlight_freestyle)
    else:
        ws.cell(row=3, column=1, value='No job data (controller offline or jobs skipped).')

    # -- 3. Plugins --------------------------------------------------------------------
    ws = wb.create_sheet('Plugins')
    ws.cell(row=1, column=1, value='Objective 3 — Installed Plugins').font = TITLE_FONT
    if pl_header:
        write_table(ws, pl_header, pl_rows, start_row=3)

    if ph_header:
        ws = wb.create_sheet('Plugin Health')
        ws.cell(row=1, column=1, value='Objective 3 — Plugin Health Report').font = TITLE_FONT
        write_table(ws, ph_header, ph_rows, start_row=3)

    wb.save(out_path)


# ======================================================================================
# Operations Center workbook
# ======================================================================================
def build_oc_workbook(oc_dir, out_path):
    info = read_json(oc_dir / 'oc-info.json') or {}
    disc_header, disc_rows = read_csv(oc_dir / 'controllers-discovered.csv')
    pl_header, pl_rows = read_csv(oc_dir / 'oc-plugins.csv')
    ph_header, ph_rows = read_csv(oc_dir / 'oc-plugin-health-report.csv')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Overview'
    ws.cell(row=1, column=1, value='Operations Center (CJOC)').font = TITLE_FONT
    kv_sheet(ws, 'Key facts', [
        ('Name', info.get('name')),
        ('URL', info.get('url')),
        ('CloudBees CI Version', info.get('version')),
        ('Connected Controllers', len(disc_rows)),
        ('Installed Plugins', len(pl_rows)),
    ], start_row=3)

    ws = wb.create_sheet('Controllers Discovered')
    ws.cell(row=1, column=1, value='Connected Controllers (Objective 1 manifest)').font = TITLE_FONT
    if disc_header:
        write_table(ws, disc_header, disc_rows, start_row=3)

    ws = wb.create_sheet('Plugins')
    ws.cell(row=1, column=1, value='Objective 3 — CJOC Installed Plugins').font = TITLE_FONT
    if pl_header:
        write_table(ws, pl_header, pl_rows, start_row=3)

    if ph_header:
        ws = wb.create_sheet('Plugin Health')
        ws.cell(row=1, column=1, value='Objective 3 — CJOC Plugin Health Report').font = TITLE_FONT
        write_table(ws, ph_header, ph_rows, start_row=3)

    wb.save(out_path)


# ======================================================================================
# Summary workbook
# ======================================================================================
def _flag_all(header, row):
    return FLAG_FILL


def build_summary_workbook(root, oc_dir, org_dir, controllers, out_path):
    wb = Workbook()

    # -- Read Me -----------------------------------------------------------------------
    ws = wb.active
    ws.title = 'Read Me'
    ws.cell(row=1, column=1, value='CloudBees CI — Environment Inventory').font = TITLE_FONT
    ws.cell(row=2, column=1, value='CIBC Success Path — Inventory & Cleanup').font = SUBTITLE_FONT
    meta = read_json(root / 'run-metadata.json') or {}
    intro = [
        ('Operations Center', meta.get('oc_url', '')),
        ('Controllers discovered', meta.get('controllers_discovered', '')),
        ('Run started', meta.get('started', '')),
        ('Run finished', meta.get('finished', '')),
    ]
    r = kv_sheet(ws, 'Run details', intro, start_row=4)
    ws.cell(row=r, column=1, value='What is in this report pack').font = SUBTITLE_FONT
    guide_header = ['Workbook / Sheet', 'What it shows', 'Objective']
    guide_rows = [
        ['00_..._Summary.xlsx', 'This workbook — environment-wide rollups and a guide', 'All'],
        ['  • Controllers', 'Every controller: CPU, memory, HA, state', '1'],
        ['  • Job Ecosystem', 'Job-type counts per controller + totals', '2'],
        ['  • Plugins Ranked', 'All plugins ranked by usage, with # controllers', '3'],
        ['  • Flagged Plugins', 'Deprecated / community / low-health / CVE plugins', '3'],
        ['  • Version Drift', 'Plugins installed at differing versions', '3'],
        ['  • Errors', 'Any per-instance collection failures', '—'],
        ['operations-center.xlsx', 'The CJOC: version, connected controllers, plugins', '1 & 3'],
        ['controller__<name>.xlsx', 'One per controller: resources, jobs, plugins', '1, 2, 3'],
    ]
    write_table(ws, guide_header, guide_rows, start_row=r + 1)

    # -- Objective rollups -------------------------------------------------------------
    def add(sheet_name, title, csv_path, highlight=None):
        h, rows = read_csv(csv_path)
        ws = wb.create_sheet(sheet_name)
        ws.cell(row=1, column=1, value=title).font = TITLE_FONT
        if h:
            write_table(ws, h, rows, start_row=3, highlight=highlight)
        else:
            ws.cell(row=3, column=1, value='(no data)')

    add('Controllers', 'Objective 1 — Controller Resource Summary', org_dir / 'controller-resource-summary.csv')
    add('Job Ecosystem', 'Objective 2 — Job Ecosystem Summary', org_dir / 'job-ecosystem-summary.csv')
    add('Plugins Ranked', 'Objective 3 — Plugins Ranked by Usage', org_dir / 'plugin-usage-ranked.csv')
    add('Flagged Plugins', 'Objective 3 — Flagged Plugins (review/replace)', org_dir / 'flagged-plugins.csv', highlight=_flag_all)
    add('Version Drift', 'Objective 3 — Plugin Version Drift', org_dir / 'plugin-version-drift.csv')
    add('Errors', 'Per-instance Collection Errors', root / 'inventory-errors.csv')

    wb.save(out_path)


# ======================================================================================
# Driver
# ======================================================================================
def sanitize(name):
    return ''.join(ch if (ch.isalnum() or ch in '-_.') else '-' for ch in str(name)).strip('-') or 'unnamed'


def latest_run(output_dir):
    runs = sorted(Path(output_dir).glob('*_environment-inventory'))
    return runs[-1] if runs else None


def parse_args():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument('--run-dir', dest='run_dir', default=None,
                   help='A specific timestamped run folder. Default: newest under ./output')
    p.add_argument('--output-dir', dest='output_dir', default=None,
                   help='Where to write the .xlsx files. Default: <run-dir>/reports')
    return p.parse_args()


def main():
    args = parse_args()
    script_dir = Path(__file__).resolve().parent
    run_dir = Path(args.run_dir) if args.run_dir else latest_run(script_dir / 'output')
    if not run_dir or not run_dir.is_dir():
        print("ERROR: no run folder found. Pass --run-dir <timestamped folder> or run the inventory first.")
        sys.exit(1)

    oc_dir = run_dir / 'operations-center'
    org_dir = run_dir / 'org-wide'
    controllers_dir = run_dir / 'controllers'
    out_dir = Path(args.output_dir) if args.output_dir else (run_dir / 'reports')
    out_dir.mkdir(parents=True, exist_ok=True)
    print(f" --> Building Excel reports from {run_dir}")
    print(f" --> Writing to {out_dir}")

    # Summary
    controllers = sorted([d for d in controllers_dir.iterdir() if d.is_dir()]) if controllers_dir.is_dir() else []
    summary_path = out_dir / '00_Environment_Inventory_Summary.xlsx'
    build_summary_workbook(run_dir, oc_dir, org_dir, controllers, summary_path)
    print(f"     + {summary_path.name}")

    # OC
    if oc_dir.is_dir():
        oc_path = out_dir / 'operations-center.xlsx'
        build_oc_workbook(oc_dir, oc_path)
        print(f"     + {oc_path.name}")

    # Per controller
    for folder in controllers:
        res = read_json(folder / 'resources.json') or {}
        name = sanitize(res.get('name') or folder.name)
        out_path = out_dir / f"controller__{name}.xlsx"
        build_controller_workbook(folder, out_path)
        print(f"     + {out_path.name}")

    print(f" --> Done. {1 + (1 if oc_dir.is_dir() else 0) + len(controllers)} workbook(s) in {out_dir}")


if __name__ == '__main__':
    main()
